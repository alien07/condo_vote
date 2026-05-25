from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.worksheet.datavalidation import DataValidation


OUTPUT = Path("public/templates/condovotes-master-import-template.xlsx")
DATA_START_ROW = 5
DATA_END_ROW = 504


def style_header(ws, headers):
    fill = PatternFill("solid", fgColor="0F766E")
    font = Font(color="FFFFFF", bold=True)
    border = Border(bottom=Side(style="thin", color="94A3B8"))

    for col_index, header in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=col_index, value=header)
        cell.fill = fill
        cell.font = font
        cell.border = border
        cell.alignment = Alignment(horizontal="center")
        cell.protection = Protection(locked=True)


def unlock_input_area(ws, column_count):
    for row in range(DATA_START_ROW, DATA_END_ROW + 1):
        for col in range(1, column_count + 1):
            ws.cell(row=row, column=col).protection = Protection(locked=False)


def add_sheet_common(ws, title, description, headers, widths):
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A5"
    ws["A1"] = title
    ws["A1"].font = Font(size=16, bold=True, color="0F172A")
    ws["A2"] = description
    ws["A2"].font = Font(color="475569")
    ws["A2"].alignment = Alignment(wrap_text=True)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
    ws.row_dimensions[2].height = 36

    style_header(ws, headers)
    unlock_input_area(ws, len(headers))

    for col_letter, width in widths.items():
        ws.column_dimensions[col_letter].width = width

    ws.auto_filter.ref = f"A4:{chr(64 + len(headers))}{DATA_END_ROW}"
    ws.protection.sheet = True
    ws.protection.enable()


def add_validations(ws, active_col):
    action_validation = DataValidation(
        type="list",
        formula1='"upsert"',
        allow_blank=False,
        showErrorMessage=True,
        error="Only upsert is allowed. Delete master data from the edit menu.",
    )
    active_validation = DataValidation(
        type="list",
        formula1='"true,false"',
        allow_blank=True,
        showErrorMessage=True,
        error="Use true or false.",
    )
    ws.add_data_validation(action_validation)
    ws.add_data_validation(active_validation)
    action_validation.add(f"A{DATA_START_ROW}:A{DATA_END_ROW}")
    active_validation.add(f"{active_col}{DATA_START_ROW}:{active_col}{DATA_END_ROW}")

    for row in range(DATA_START_ROW, DATA_END_ROW + 1):
        ws.cell(row=row, column=1, value="upsert")


def main():
    wb = Workbook()
    instructions = wb.active
    instructions.title = "Instructions"
    instructions.sheet_view.showGridLines = False
    instructions["A1"] = "condoVotes master data import template"
    instructions["A1"].font = Font(size=16, bold=True, color="0F172A")
    instructions["A3"] = "Workflow"
    instructions["A3"].font = Font(bold=True)
    instruction_rows = [
        "1. Fill only the unlocked rows in Rooms and Owners.",
        "2. Keep import_action as upsert. This template intentionally does not support delete.",
        "3. Upload this same .xlsx file from Admin > People.",
        "4. Rooms are upserted by room_number.",
        "5. Owners are upserted by email.",
        "6. To delete or deactivate master data, use the update/edit menu in the app.",
    ]
    for index, text in enumerate(instruction_rows, start=4):
        instructions.cell(row=index, column=1, value=text)
    instructions.column_dimensions["A"].width = 110
    instructions.protection.sheet = True
    instructions.protection.enable()

    rooms = wb.create_sheet("Rooms")
    room_headers = [
        "import_action",
        "room_number",
        "ownership_percent",
        "building",
        "floor",
        "area_size",
        "active",
    ]
    add_sheet_common(
        rooms,
        "Rooms",
        "Required: room_number, ownership_percent. Optional: building, floor, area_size, active.",
        room_headers,
        {
            "A": 16,
            "B": 18,
            "C": 20,
            "D": 16,
            "E": 12,
            "F": 14,
            "G": 12,
        },
    )
    add_validations(rooms, "G")

    owners = wb.create_sheet("Owners")
    owner_headers = [
        "import_action",
        "full_name",
        "email",
        "phone",
        "line_id",
        "active",
    ]
    add_sheet_common(
        owners,
        "Owners",
        "Required: full_name, email. Optional: phone, line_id, active.",
        owner_headers,
        {
            "A": 16,
            "B": 28,
            "C": 32,
            "D": 18,
            "E": 18,
            "F": 12,
        },
    )
    add_validations(owners, "F")
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT)


if __name__ == "__main__":
    main()
